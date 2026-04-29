"""Excel output renderer matching the CMA spreadsheet layout.

Generates a populated workbook with the same tab structure as the
CMA Lift Station Calculations Template:

    Project Data | Flow Calculations | Flow Rate | Pump Cycle Times |
    Headloss Fitting | TDH | System Curves | Buoyancy | K-Value

Per the xlsx skill, formulas are used where possible so the recipient can
edit assumptions and have downstream cells recompute. Hardcoded values are
limited to inputs.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..calc.design import DesignReport
from ..catalog import k_values

# Style constants
_HEADER_FILL = PatternFill("solid", start_color="305496")
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
_LABEL_FONT = Font(name="Arial", bold=True, size=10)
_INPUT_FONT = Font(name="Arial", color="0000FF", size=10)  # blue for inputs
_FORMULA_FONT = Font(name="Arial", color="000000", size=10)
_NORMAL_FONT = Font(name="Arial", size=10)
_PLACEHOLDER_FILL = PatternFill("solid", start_color="FFFF00")
_THIN = Side(border_style="thin", color="000000")
_BOX = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)


def _set_column_widths(ws, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _write_project_data(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Project Data")
    ws["A1"] = "PROJECT DATA"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:D1")

    rows = [
        ("Project Name:", report.project.project_name, "Designed:", report.project.designer),
        ("Project Location:", report.project.project_location, "Reviewed:", ""),
        ("Project Number:", report.project.project_number, "Checked:", ""),
        ("Date:", "", "", ""),
    ]
    for r, (a, b, c, d) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=a).font = _LABEL_FONT
        ws.cell(row=r, column=2, value=b).font = _INPUT_FONT
        ws.cell(row=r, column=3, value=c).font = _LABEL_FONT
        ws.cell(row=r, column=4, value=d).font = _INPUT_FONT

    _set_column_widths(ws, {"A": 22, "B": 30, "C": 14, "D": 22})


def _write_flow_calculations(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Flow Calculations")
    ws["A1"] = "FLOW CALCULATIONS (FAC 64E-6.008 Table I)"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:D1")

    headers = ["Type of Structure / Land Use", "Number of Units",
               "Flow per Unit (GPD)", "Average Daily Flow (GPD)"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = _LABEL_FONT
        cell.border = _BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center")

    if report.flow_estimate is not None:
        first_data = 4
        for i, item in enumerate(report.flow_estimate.line_items):
            r = first_data + i
            ws.cell(row=r, column=1, value=item.description).font = _INPUT_FONT
            ws.cell(row=r, column=2, value=item.number_of_units).font = _INPUT_FONT
            ws.cell(row=r, column=3, value=item.flow_per_unit_gpd).font = _INPUT_FONT
            ws.cell(row=r, column=4, value=f"=B{r}*C{r}").font = _FORMULA_FONT
        last_data = first_data + len(report.flow_estimate.line_items) - 1

        # Total row
        total_row = last_data + 2
        ws.cell(row=total_row, column=1, value="TOTAL ADF (GPD)").font = _LABEL_FONT
        ws.cell(
            row=total_row,
            column=4,
            value=f"=SUM(D{first_data}:D{last_data})",
        ).font = _FORMULA_FONT

        ws.cell(row=total_row + 1, column=1, value="Service Area Population").font = _LABEL_FONT
        ws.cell(row=total_row + 1, column=4,
                value=report.flow_estimate.service_area_population).font = _INPUT_FONT

    _set_column_widths(ws, {"A": 50, "B": 16, "C": 18, "D": 22})


def _write_flow_rate(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Flow Rate")
    ws["A1"] = "DESIGN PUMPING FLOW RATE"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:D1")

    if report.pumping_flow is None:
        return
    pf = report.pumping_flow

    ws.cell(row=3, column=1, value="Average Daily Flow (GPD):").font = _LABEL_FONT
    ws.cell(row=3, column=2, value=pf.average_daily_flow_gpd).font = _INPUT_FONT

    ws.cell(row=4, column=1, value="Service Area Population:").font = _LABEL_FONT
    pop = (
        report.flow_estimate.service_area_population
        if report.flow_estimate is not None
        else 0.0
    )
    ws.cell(row=4, column=2, value=pop).font = _INPUT_FONT

    ws.cell(row=6, column=1, value="--- Peak Factor (both methods, take greater) ---").font = _LABEL_FONT

    ws.cell(row=7, column=1, value="Tiered Table PF:").font = _LABEL_FONT
    ws.cell(row=7, column=2, value=pf.peak_factor_result.table_pf).font = _FORMULA_FONT

    ws.cell(row=8, column=1, value="Fair-Geyer PF (P in thousands):").font = _LABEL_FONT
    ws.cell(row=8, column=2,
            value=f"=MAX(2.5,(18+SQRT(B4/1000))/(4+SQRT(B4/1000)))").font = _FORMULA_FONT

    ws.cell(row=9, column=1, value="Design Peak Factor:").font = _LABEL_FONT
    ws.cell(row=9, column=2, value="=MAX(B7,B8)").font = _FORMULA_FONT

    ws.cell(row=10, column=1, value="Governing Method:").font = _LABEL_FONT
    ws.cell(row=10, column=2, value=pf.peak_factor_result.governing_method).font = _NORMAL_FONT

    ws.cell(row=12, column=1, value="Peak Flow (GPD):").font = _LABEL_FONT
    ws.cell(row=12, column=2, value="=B3*B9").font = _FORMULA_FONT

    ws.cell(row=13, column=1, value="Min Pumping Rate (GPM):").font = _LABEL_FONT
    ws.cell(row=13, column=2, value="=B12/(24*60)").font = _FORMULA_FONT

    ws.cell(row=14, column=1, value="Design Pumping Rate (GPM, round up to 100):").font = _LABEL_FONT
    ws.cell(row=14, column=2, value="=ROUNDUP(B13,-2)").font = _FORMULA_FONT

    _set_column_widths(ws, {"A": 50, "B": 18})


def _write_pump_cycle_times(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Pump Cycle Times")
    ws["A1"] = "PUMP CYCLE TIMES"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:E1")

    g = report.project.geometry
    ws.cell(row=3, column=1, value="Station Number").font = _LABEL_FONT
    ws.cell(row=3, column=2, value=g.station_number).font = _INPUT_FONT
    ws.cell(row=4, column=1, value="Wet Well Diameter (ft)").font = _LABEL_FONT
    ws.cell(row=4, column=2, value=g.diameter_ft).font = _INPUT_FONT
    ws.cell(row=5, column=1, value="100-YR Flood Elevation").font = _LABEL_FONT
    ws.cell(row=5, column=2, value=str(g.flood_elevation)).font = _INPUT_FONT

    elev_rows = [
        ("Top Slab Elevation - EL A", g.top_slab_el),
        ("Finish Grade - EL B", g.finish_grade_el),
        ("Centerline Discharge - EL C", g.centerline_discharge_pipe_el),
        ("Lowest Influent Line - EL E", g.lowest_influent_line_el),
        ("Lead Pump On - EL F", g.lead_pump_on_el),
        ("Both Pumps Off - EL G", g.both_pumps_off_el),
        ("Pump Suction - EL H", g.pump_suction_el),
        ("Wet Well Base Slab - EL I", g.wetwell_base_slab_el),
        ("Pump Discharge Diameter (in)", g.pump_discharge_diameter_in),
        ("Discharge Pipe Diameter (in)", g.discharge_pipe_diameter_in),
        ("Common FM Diameter (in)", g.common_force_main_diameter_in),
        ("Pump Dimension (in)", g.pump_dimension_in),
    ]
    for i, (label, val) in enumerate(elev_rows, start=7):
        ws.cell(row=i, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=i, column=2, value=val).font = _INPUT_FONT
        ws.cell(row=i, column=3, value="FT-NAVD" if "El" in label or "EL" in label else "").font = _NORMAL_FONT

    # Wet well volume formula:
    vol_row = 21
    ws.cell(row=vol_row, column=1, value="Working Volume (gal)").font = _LABEL_FONT
    # V = pi * D^2 / 4 * (EL_F - EL_G) * 7.48
    # B4=diameter, B11=EL F, B12=EL G
    ws.cell(row=vol_row, column=2, value="=PI()*B4^2/4*(B11-B12)*7.48").font = _FORMULA_FONT
    ws.cell(row=vol_row, column=3, value="GAL").font = _NORMAL_FONT

    if report.cycle_times is not None:
        ct = report.cycle_times
        cyc_start = vol_row + 3
        ws.cell(row=cyc_start, column=1, value="Q avg (GPM)").font = _LABEL_FONT
        ws.cell(row=cyc_start, column=2, value=ct.q_avg_gpm).font = _FORMULA_FONT
        ws.cell(row=cyc_start + 1, column=1, value="Q peak (GPM)").font = _LABEL_FONT
        ws.cell(row=cyc_start + 1, column=2, value=ct.q_peak_gpm).font = _FORMULA_FONT
        ws.cell(row=cyc_start + 2, column=1, value="Q pump (GPM)").font = _LABEL_FONT
        ws.cell(row=cyc_start + 2, column=2, value=ct.q_pump_gpm).font = _INPUT_FONT

        v_ref = f"B{vol_row}"
        qa, qp, qpump = (
            f"B{cyc_start}",
            f"B{cyc_start + 1}",
            f"B{cyc_start + 2}",
        )
        rows = [
            ("Tmin = 4*V/Qpump", f"=4*{v_ref}/{qpump}", "min (>5)"),
            ("Tavg run = V/(Qpeak-Qavg)", f"={v_ref}/({qp}-{qa})", "min (>2)"),
            ("Tavg cycle = Tavg run + V/Qavg",
             f"={v_ref}/({qp}-{qa})+{v_ref}/{qa}", "min (<30)"),
            ("Tpeak run = V/(Qpump-Qpeak)", f"={v_ref}/({qpump}-{qp})", "min"),
            ("Tpeak cycle = Tpeak run + V/Qpeak",
             f"={v_ref}/({qpump}-{qp})+{v_ref}/{qp}", "min"),
        ]
        for i, (label, formula, unit) in enumerate(rows, start=cyc_start + 4):
            ws.cell(row=i, column=1, value=label).font = _LABEL_FONT
            ws.cell(row=i, column=2, value=formula).font = _FORMULA_FONT
            ws.cell(row=i, column=3, value=unit).font = _NORMAL_FONT

        ws.cell(row=cyc_start + 10, column=1,
                value="DIVERGENCE NOTE: Tavg uses (Qpeak-Qavg), not (Qpump-Qavg), per CMA spreadsheet."
                ).font = Font(name="Arial", italic=True, color="800000", size=9)

    _set_column_widths(ws, {"A": 45, "B": 18, "C": 12})


def _write_headloss_fitting(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Headloss Fitting")
    ws["A1"] = "TABULATE HEADLOSSES DUE TO FITTINGS"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:E1")

    sh = report.static_head
    if sh is not None:
        ws.cell(row=3, column=1, value="Destination Elevation (ft)").font = _LABEL_FONT
        ws.cell(row=3, column=2, value=sh.destination_elevation_ft).font = _INPUT_FONT
        ws.cell(row=4, column=1, value="Pump Off Elevation (ft)").font = _LABEL_FONT
        ws.cell(row=4, column=2, value=sh.pump_off_elevation_ft).font = _INPUT_FONT
        ws.cell(row=5, column=1, value="Static Lift = Dest - Pump Off").font = _LABEL_FONT
        ws.cell(row=5, column=2, value="=B3-B4").font = _FORMULA_FONT
        ws.cell(row=6, column=1, value="Static Head at Connection (psi)").font = _LABEL_FONT
        ws.cell(row=6, column=2, value=sh.static_head_at_connection_psi).font = _INPUT_FONT
        ws.cell(row=7, column=1, value="Static Head at Connection (ft)").font = _LABEL_FONT
        ws.cell(row=7, column=2, value="=B6*2.31").font = _FORMULA_FONT
        ws.cell(row=8, column=1, value="Total Static Head (ft)").font = _LABEL_FONT
        ws.cell(row=8, column=2, value="=B5+B7").font = _FORMULA_FONT

    # Force main segment fittings
    row = 11
    for seg in report.project.force_main:
        ws.cell(row=row, column=1, value=f"FM Segment {seg.label}").font = _HEADER_FONT
        ws.cell(row=row, column=1).fill = _HEADER_FILL
        ws.cell(row=row, column=1).font = Font(name="Arial", bold=True, color="FFFFFF")
        row += 1
        ws.cell(row=row, column=1, value="Diameter (in)").font = _LABEL_FONT
        ws.cell(row=row, column=2, value=seg.diameter_in).font = _INPUT_FONT
        row += 1
        ws.cell(row=row, column=1, value="Length (ft)").font = _LABEL_FONT
        ws.cell(row=row, column=2, value=seg.length_ft).font = _INPUT_FONT
        row += 1
        ws.cell(row=row, column=1, value="C Factor").font = _LABEL_FONT
        ws.cell(row=row, column=2, value=seg.c_factor).font = _INPUT_FONT
        row += 2
        ws.cell(row=row, column=1, value="Fitting").font = _LABEL_FONT
        ws.cell(row=row, column=2, value="Quantity").font = _LABEL_FONT
        ws.cell(row=row, column=3, value="K Value").font = _LABEL_FONT
        ws.cell(row=row, column=4, value="Total K").font = _LABEL_FONT
        row += 1
        first_fit = row
        for fit in seg.fittings:
            ws.cell(row=row, column=1, value=fit.name).font = _INPUT_FONT
            ws.cell(row=row, column=2, value=fit.quantity).font = _INPUT_FONT
            ws.cell(row=row, column=3, value=fit.k_value).font = _INPUT_FONT
            ws.cell(row=row, column=4, value=f"=B{row}*C{row}").font = _FORMULA_FONT
            row += 1
        if seg.fittings:
            ws.cell(row=row, column=3, value="Total K:").font = _LABEL_FONT
            ws.cell(row=row, column=4, value=f"=SUM(D{first_fit}:D{row - 1})").font = _FORMULA_FONT
            row += 2
        else:
            row += 1

    _set_column_widths(ws, {"A": 32, "B": 14, "C": 14, "D": 14})


def _write_tdh(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("TDH")
    ws["A1"] = "TOTAL DYNAMIC HEAD"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:E1")

    if report.tdh_curve is None:
        return

    headers = ["Flow (GPM)", "Static Head (ft)"]
    for seg in report.tdh_curve.segments:
        headers.extend([
            f"Seg {seg.label} Velocity (ft/s)",
            f"Seg {seg.label} Friction (ft)",
            f"Seg {seg.label} Fittings (ft)",
        ])
    headers.append("TDH (ft)")
    if report.selected_pump is not None:
        headers.append(f"Pump Head: {report.selected_pump.model} (ft)")

    for c, h in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=h).font = _LABEL_FONT
        ws.cell(row=3, column=c).alignment = Alignment(
            wrap_text=True, horizontal="center"
        )

    for i, point in enumerate(report.tdh_curve.points, start=4):
        ws.cell(row=i, column=1, value=point.flow_gpm)
        ws.cell(row=i, column=2, value=point.static_head_ft)
        c = 3
        for seg in point.segment_losses:
            ws.cell(row=i, column=c, value=seg.velocity_ft_s)
            ws.cell(row=i, column=c + 1, value=seg.friction_loss_ft)
            ws.cell(row=i, column=c + 2, value=seg.fitting_loss_ft)
            c += 3
        ws.cell(row=i, column=c, value=point.total_dynamic_head_ft)
        if report.selected_pump is not None:
            ph = report.selected_pump.head_at(point.flow_gpm)
            ws.cell(row=i, column=c + 1, value=ph if ph == ph else None)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def _write_system_curves(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("System Curves")
    ws["A1"] = "LIFT STATION SYSTEM AND PUMP CURVES"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:C1")

    if report.tdh_curve is None:
        return

    ws["A3"] = "Flow (GPM)"
    ws["B3"] = "System Curve TDH (ft)"
    if report.selected_pump is not None:
        ws["C3"] = f"Pump Curve: {report.selected_pump.model} (ft)"

    n_points = len(report.tdh_curve.points)
    for i, p in enumerate(report.tdh_curve.points, start=4):
        ws.cell(row=i, column=1, value=p.flow_gpm)
        ws.cell(row=i, column=2, value=p.total_dynamic_head_ft)
        if report.selected_pump is not None:
            ph = report.selected_pump.head_at(p.flow_gpm)
            ws.cell(row=i, column=3, value=ph if ph == ph else None)

    chart = LineChart()
    chart.title = "System Curve and Pump Curve"
    chart.y_axis.title = "Total Dynamic Head (ft)"
    chart.x_axis.title = "Flow (GPM)"
    chart.height = 10
    chart.width = 18
    last_row = 3 + n_points

    sys_data = Reference(
        ws, min_col=2, min_row=3, max_col=2, max_row=last_row
    )
    chart.add_data(sys_data, titles_from_data=True)

    if report.selected_pump is not None:
        pump_data = Reference(
            ws, min_col=3, min_row=3, max_col=3, max_row=last_row
        )
        chart.add_data(pump_data, titles_from_data=True)

    flows_ref = Reference(ws, min_col=1, min_row=4, max_row=last_row)
    chart.set_categories(flows_ref)
    ws.add_chart(chart, "E3")

    _set_column_widths(ws, {"A": 14, "B": 22, "C": 28})


def _write_buoyancy(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Buoyancy (Tremie)")
    ws["A1"] = "BUOYANCY CALCULATION FOR WET WELL (single-case)"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:D1")

    if report.buoyancy is None:
        ws.cell(row=3, column=1, value="No buoyancy inputs provided.").font = _NORMAL_FONT
        return

    b = report.buoyancy
    rows = [
        ("Outside Diameter (ft)", b.outside_diameter_ft),
        ("Height of Structure (ft)", b.height_of_structure_ft),
        ("Tremie Seal Height (ft)", b.tremie_seal_height_ft),
        ("Volume Displaced (CF)", b.volume_displaced_cf),
        ("Buoyant Force (lb)", b.buoyant_force_lb),
        ("Factored Buoyant Force (lb)", b.factored_buoyant_force_lb),
        ("Volume of Walls (CF)", b.volume_of_walls_cf),
        ("Volume of Tremie Seal (CF)", b.volume_of_tremie_seal_cf),
        ("Total Resisting Weight (lb)", b.total_weight_lb),
        ("Weight / Buoyancy Ratio", b.weight_to_buoyancy_ratio),
        ("Pass Check?", "OK" if b.passes_check else "NOT OK"),
    ]
    for i, (label, val) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=i, column=2, value=val).font = _NORMAL_FONT

    note_row = 3 + len(rows) + 1
    ws.cell(row=note_row, column=1,
            value="DIVERGENCE: This implements the spreadsheet's single-case "
            "calculation, not the two-load-case CMA standard (Notes 22 & 23)."
            ).font = Font(name="Arial", italic=True, color="800000")

    _set_column_widths(ws, {"A": 40, "B": 18})


def _write_k_values(wb: Workbook) -> None:
    ws = wb.create_sheet("K-Value")
    ws["A1"] = "K-VALUE REFERENCE TABLE (Crane TP-410 style)"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL

    buckets = k_values.pipe_size_buckets()
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(buckets) + 2)

    ws.cell(row=3, column=1, value="Fitting").font = _LABEL_FONT
    for i, (_, label) in enumerate(buckets, start=2):
        ws.cell(row=3, column=i, value=label).font = _LABEL_FONT
        ws.cell(row=3, column=i).alignment = Alignment(horizontal="center")
    for i, row in enumerate(k_values.all_rows(), start=4):
        ws.cell(row=i, column=1, value=row.name).font = _NORMAL_FONT
        for j, k in enumerate(row.k_values, start=2):
            ws.cell(row=i, column=j, value=k).font = _NORMAL_FONT

    ws.column_dimensions["A"].width = 28
    for j in range(2, len(buckets) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 9


def _write_notes(wb: Workbook, report: DesignReport) -> None:
    ws = wb.create_sheet("Design Notes")
    ws["A1"] = "DESIGN NOTES & WARNINGS"
    ws["A1"].font = _HEADER_FONT
    ws["A1"].fill = _HEADER_FILL
    if not report.notes:
        ws["A3"] = "No notes."
        return
    for i, note in enumerate(report.notes, start=3):
        ws.cell(row=i, column=1, value=f"- {note}").font = _NORMAL_FONT
    ws.column_dimensions["A"].width = 100


def render_excel(report: DesignReport, path: Path | str) -> Path:
    """Render a DesignReport as a .xlsx workbook.

    Tab structure mirrors the CMA Lift Station Calculations Template.
    """
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    # Drop the default empty sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    _write_project_data(wb, report)
    _write_flow_calculations(wb, report)
    _write_flow_rate(wb, report)
    _write_pump_cycle_times(wb, report)
    _write_headloss_fitting(wb, report)
    _write_tdh(wb, report)
    _write_system_curves(wb, report)
    _write_buoyancy(wb, report)
    _write_k_values(wb)
    _write_notes(wb, report)

    wb.save(p)
    return p
